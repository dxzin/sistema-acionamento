import openpyxl

wb = openpyxl.load_workbook('referencias/TESTE.xlsx', data_only=False)

def inspect_aux_columns():
    ws = wb['02.06.26  - UBJP']
    print("=== 02.06.26  - UBJP: Columns S to V ===")
    for r in range(1, 10):
        print(f"Row {r:2d}: S={ws.cell(r, 19).value}, T={ws.cell(r, 20).value}, U={ws.cell(r, 21).value}, V={ws.cell(r, 22).value}, AB={ws.cell(r, 28).value}")

    ws_aug = wb['01.08.26 - UBJP']
    print("\n=== 01.08.26 - UBJP: Columns V to Y ===")
    for r in range(1, 10):
        print(f"Row {r:2d}: V={ws_aug.cell(r, 22).value}, W={ws_aug.cell(r, 23).value}, X={ws_aug.cell(r, 24).value}, Y={ws_aug.cell(r, 25).value}")

inspect_aux_columns()
