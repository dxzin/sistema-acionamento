"""Sincroniza o cadastro de condutores com as escalas de setembro de 2026.

Por padrão executa apenas a validação. Use --apply para gravar no Supabase.
Os cadastros anteriores são inativados, nunca excluídos, preservando as FKs históricas.
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "referencias" / "TESTE.xlsx"
SHEETS = ("Escala Setembro 2026 - SP", "Escala Setembro 2026 - MG")
HISTORY_SHEETS = (
    "Escala Junho 2026 - SP", "Escala Junho 2026 - MG",
    "Escala Julho 2026 - SP", "Escala Julho 2026 - MG",
    "Escala Agosto 2026 - SP", "Escala Agosto 2026 - MG",
)


def normalize_name(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def match_key(value: object) -> str:
    text = unicodedata.normalize("NFKD", normalize_name(value)).encode("ascii", "ignore").decode().upper()
    return re.sub(r"[^A-Z0-9]", "", text)


def clean_login(value: object) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def parse_time(value: object) -> str | None:
    if isinstance(value, dt.time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, (int, float)) and 0 <= value < 1:
        seconds = round(value * 86400) % 86400
        return f"{seconds // 3600:02d}:{seconds % 3600 // 60:02d}:00"
    if isinstance(value, str) and re.fullmatch(r"\d{1,2}:\d{2}(?::\d{2})?", value.strip()):
        parts = value.strip().split(":")
        return f"{int(parts[0]):02d}:{int(parts[1]):02d}:{int(parts[2]) if len(parts) > 2 else 0:02d}"
    return None


def row_metadata(values: tuple[object, ...], first_meta_column: int) -> tuple[str | None, str | None]:
    metadata = values[first_meta_column - 1:]
    start_time = next((parsed for value in metadata if (parsed := parse_time(value))), None)
    shift = next((str(value).strip().upper() for value in metadata if isinstance(value, str) and re.fullmatch(r"[A-Z]", value.strip().upper())), None)
    return shift, start_time


def load_env() -> dict[str, str]:
    result: dict[str, str] = {}
    for line in (ROOT / ".env.local").read_text(encoding="utf-8").splitlines():
        if line and not line.lstrip().startswith("#") and "=" in line:
            key, value = line.split("=", 1)
            result[key.strip()] = value.strip().strip('"').strip("'")
    return result


def request_json(url: str, key: str, method: str = "GET", body: object | None = None, prefer: str | None = None):
    headers = {"apikey": key, "Authorization": f"Bearer {key}"}
    data = None
    if body is not None:
        headers["Content-Type"] = "application/json"
        data = json.dumps(body, ensure_ascii=False).encode("utf-8")
    if prefer:
        headers["Prefer"] = prefer
    request = urllib.request.Request(url, data=data, headers=headers, method=method)
    try:
        with urllib.request.urlopen(request, timeout=60) as response:
            raw = response.read()
            return json.loads(raw) if raw else None
    except urllib.error.HTTPError as error:
        detail = error.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"Supabase retornou HTTP {error.code}: {detail}") from error


def extract_roster():
    workbook = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    prior_logins: dict[str, list[str]] = defaultdict(list)
    prior_metadata: dict[str, tuple[str | None, str | None]] = {}

    for sheet_name in HISTORY_SHEETS:
        sheet = workbook[sheet_name]
        is_mg = sheet_name.endswith("MG")
        login_index = 2 if is_mg else 1
        day_count = 31 if "Julho" in sheet_name or "Agosto" in sheet_name else 30
        first_meta_column = (4 if is_mg else 3) + day_count
        for values in sheet.iter_rows(min_row=2, values_only=True):
            if not values[0]:
                continue
            key = match_key(values[0])
            login = clean_login(values[login_index])
            if login:
                prior_logins[key].append(login)
            metadata = row_metadata(values, first_meta_column)
            if any(metadata):
                prior_metadata[key] = metadata

    roster = []
    source_logins: Counter[str] = Counter()
    for sheet_name in SHEETS:
        sheet = workbook[sheet_name]
        is_mg = sheet_name.endswith("MG")
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            if not values[0]:
                continue
            original_name = normalize_name(values[0])
            name = re.sub(r"\s+-\s+VOLTA\s+\d{1,2}/\d{1,2}\s*$", "", original_name, flags=re.IGNORECASE)
            key = match_key(name)
            login = clean_login(values[2 if is_mg else 1])
            if login:
                source_logins[login] += 1
            if is_mg:
                base_code = normalize_name(values[1]).upper()
                shift, start_time = row_metadata(values, 34)
            else:
                shift, start_time = row_metadata(values, 34)
                base_code = "LP" if parse_time(values[33]) else "CASTILHO"

            roster.append({
                "source": f"{sheet_name}!{row_number}",
                "name": name,
                "key": key,
                "login": login,
                "base_code": base_code,
                "default_shift": shift,
                "default_start_time": start_time,
            })

    used_logins: set[str] = set()
    corrections = []
    for driver in roster:
        login = driver["login"]
        needs_fallback = not login or source_logins[login] > 1 or login in used_logins
        if needs_fallback:
            candidates = list(reversed(prior_logins.get(driver["key"], [])))
            replacement = next((candidate for candidate in candidates if candidate not in used_logins), None)
            if not replacement:
                raise ValueError(f"Sem login único para {driver['name']} ({driver['source']}).")
            corrections.append((driver["name"], login, replacement))
            driver["login"] = replacement
        used_logins.add(driver["login"])

        if not driver["default_shift"] or not driver["default_start_time"]:
            old_shift, old_time = prior_metadata.get(driver["key"], (None, None))
            driver["default_shift"] = driver["default_shift"] or old_shift
            driver["default_start_time"] = driver["default_start_time"] or old_time

    duplicates = [login for login, count in Counter(driver["login"] for driver in roster).items() if count > 1]
    invalid_bases = [driver for driver in roster if driver["base_code"] not in {"LP", "CASTILHO", "UB", "LZ", "JP"}]
    if duplicates or invalid_bases:
        raise ValueError(f"Validação falhou: duplicidades={duplicates}, bases={invalid_bases}")
    return roster, corrections


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--apply", action="store_true", help="Grava os dados validados no Supabase.")
    args = parser.parse_args()
    roster, corrections = extract_roster()
    counts = Counter(driver["base_code"] for driver in roster)
    print(json.dumps({"drivers": len(roster), "by_base": counts, "login_corrections": corrections}, ensure_ascii=False, indent=2))
    if not args.apply:
        print("Validação concluída. Nenhuma alteração foi gravada.")
        return

    env = load_env()
    base_url = env["NEXT_PUBLIC_SUPABASE_URL"].rstrip("/") + "/rest/v1"
    service_key = env["SUPABASE_SERVICE_ROLE_KEY"]
    bases = request_json(f"{base_url}/bases?select=id,code", service_key)
    base_ids = {base["code"]: base["id"] for base in bases}
    missing_bases = sorted(set(counts) - set(base_ids))
    if missing_bases:
        raise ValueError(f"Bases ausentes no Supabase: {missing_bases}")

    previous = request_json(f"{base_url}/drivers?select=id,login_matricula,is_active", service_key)
    payload = [{
        "name": driver["name"],
        "login_matricula": driver["login"],
        "base_id": base_ids[driver["base_code"]],
        "default_shift": driver["default_shift"],
        "default_start_time": driver["default_start_time"],
        "is_active": True,
    } for driver in roster]
    request_json(
        f"{base_url}/drivers?on_conflict=login_matricula",
        service_key,
        method="POST",
        body=payload,
        prefer="resolution=merge-duplicates,return=representation",
    )

    real_logins = {driver["login"] for driver in roster}
    obsolete_ids = [driver["id"] for driver in previous if driver["is_active"] and driver["login_matricula"] not in real_logins]
    if obsolete_ids:
        encoded_ids = ",".join(obsolete_ids)
        request_json(
            f"{base_url}/drivers?id=in.({urllib.parse.quote(encoded_ids, safe=',-')})",
            service_key,
            method="PATCH",
            body={"is_active": False},
            prefer="return=representation",
        )

    active = request_json(f"{base_url}/drivers?select=id,name,login_matricula,base_id&is_active=eq.true", service_key)
    if len(active) != len(roster) or len({driver['login_matricula'] for driver in active}) != len(roster):
        raise RuntimeError(f"Conferência final falhou: esperados {len(roster)}, encontrados {len(active)}.")
    print(json.dumps({"inserted_or_updated": len(roster), "inactivated": len(obsolete_ids), "active_after": len(active)}, indent=2))


if __name__ == "__main__":
    main()
