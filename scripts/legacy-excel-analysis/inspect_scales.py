import openpyxl, json

wb = openpyxl.load_workbook('referencias/TESTE.xlsx', data_only=False)

def inspect_aux_and_scales():
    print("=== 1. PLANILHA1 & GRAFICO3 ===")
    if 'Planilha1' in wb.sheetnames:
        ws = wb['Planilha1']
        print(f"Planilha1: max_row={ws.max_row}, max_col={ws.max_column}")
        for r in range(1, ws.max_row+1):
            for c in range(1, ws.max_column+1):
                if ws.cell(r, c).value is not None:
                    print(f"  P1 [{r},{c}]: {ws.cell(r, c).value}")
    
    print("\n=== 2. ESCALA MENSAL SHEETS & STATUS CODES ===")
    scale_sheets = [s for s in wb.sheetnames if 'Escala' in s]
    all_codes = set()
    for s_name in scale_sheets:
        ws = wb[s_name]
        print(f"\nScale: {s_name} | Rows: {ws.max_row}, Cols: {ws.max_column}")
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column+1)]
        print(f"  Header row 1: {headers}")
        
        # Collect unique codes in day columns (columns 3 to 33 approx)
        sheet_codes = set()
        for r in range(2, ws.max_row+1):
            for c in range(3, ws.max_column+1):
                v = ws.cell(r, c).value
                if v is not None:
                    sheet_codes.add(str(v).strip())
        print(f"  Unique codes in {s_name}: {sorted(list(sheet_codes))}")
        all_codes.update(sheet_codes)

    print(f"\n>>> ALL UNIQUE CODES IN ALL SCALES: {sorted(list(all_codes))}")

inspect_aux_and_scales()
