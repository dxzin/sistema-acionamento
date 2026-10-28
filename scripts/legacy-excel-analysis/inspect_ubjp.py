import openpyxl, json, re

wb = openpyxl.load_workbook('referencias/TESTE.xlsx', data_only=False)

def inspect_ubjp_deep():
    print("=================================================================")
    print("--- DEEP DIVE: DAILY UBJP SHEETS ---")
    for s_name in ['01.06.26  - UBJP', '02.06.26  - UBJP', '15.06.26  - UBJP', '01.07.26 - UBJP', '01.08.26 - UBJP']:
        if s_name not in wb.sheetnames:
            # try finding with strip
            matched = [s for s in wb.sheetnames if s_name.replace(' ', '') in s.replace(' ', '')]
            if matched:
                s_name = matched[0]
            else:
                continue
        ws = wb[s_name]
        print(f"\n>>> SHEET: {repr(s_name)} | Rows: {ws.max_row}, Cols: {ws.max_column}")
        for c in range(1, ws.max_column+1):
            col_letter = openpyxl.utils.get_column_letter(c)
            h4 = ws.cell(4, c).value
            v5 = ws.cell(5, c).value
            v6 = ws.cell(6, c).value
            print(f"  Col {col_letter:2s} ({c:2d}) | Header: {str(h4):22s} | R5: {str(v5)[:50]} | R6: {str(v6)[:50]}")

inspect_ubjp_deep()
