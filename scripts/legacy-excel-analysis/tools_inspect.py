import openpyxl, re, json

wb = openpyxl.load_workbook('referencias/TESTE.xlsx', data_only=False)

def inspect_daily_general_deep():
    print('=================================================================')
    print('--- DEEP DIVE: DAILY GENERAL SHEETS ---')
    for s_name in ['01.06.26', '02.06.26', '03.06.26', '01.07.26', '01.08.26']:
        if s_name not in wb.sheetnames:
            continue
        ws = wb[s_name]
        print(f'\n>>> SHEET: {s_name} | Rows: {ws.max_row}, Cols: {ws.max_column}')
        headers = [ws.cell(4, c).value for c in range(1, ws.max_column+1)]
        for idx, h in enumerate(headers, 1):
            col_letter = openpyxl.utils.get_column_letter(idx)
            val5 = ws.cell(5, idx).value
            val6 = ws.cell(6, idx).value
            f5 = str(val5) if val5 is not None else ''
            f5 = re.sub(r'Tabela\w+', '[Table]', f5)
            f6 = str(val6) if val6 is not None else ''
            f6 = re.sub(r'Tabela\w+', '[Table]', f6)
            print(f'  Col {col_letter:2s} ({idx:2d}) | {str(h):22s} | R5: {f5[:80]}')
            if f6 != f5 and f6 != '':
                print(f'         {" "*22} | R6: {f6[:80]}')

inspect_daily_general_deep()
