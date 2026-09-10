"""Sincroniza todas as folhas diárias e despachos históricos com o Supabase.

Lê as abas diárias (DD.MM.YY e DD.MM.YY - UBJP) do Excel TESTE.xlsx
e grava em:
- daily_sheets
- dispatches
- locations (cria novos locais sob demanda)
- vehicles (cria novas frotas sob demanda)
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "referencias" / "TESTE.xlsx"


def normalize_name(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def match_key(value: object) -> str:
    text = unicodedata.normalize("NFKD", normalize_name(value)).encode("ascii", "ignore").decode().upper()
    return re.sub(r"[^A-Z0-9]", "", text)


def clean_str(value: object) -> str | None:
    if value is None:
        return None
    val = str(value).strip()
    if not val or val in ("#N/A", "#VALOR!", "#REF!", "#VALUE!", "None", "SEM ID"):
        return None
    return val


def normalize_status(raw_val: object) -> str:
    if not raw_val:
        return "ACIONADO"
    val = unicodedata.normalize("NFKD", str(raw_val)).encode("ascii", "ignore").decode().upper().strip()
    if "ACIONA" in val: return "ACIONADO"
    if "ABONA" in val: return "ABONADO"
    if "APOIO" in val: return "APOIO"
    if "TROCA" in val or "BALDEIO" in val: return "TROCA"
    if "FALTA" in val: return "FALTA"
    if "ATESTA" in val: return "ATESTADO"
    if "AFASTA" in val: return "AFASTADO"
    if "BATE" in val: return "BATE E VOLTA"
    if "SUSPEN" in val: return "SUSPENSÃO"
    if "FECHADA" in val: return "VIAGEM FECHADA"
    if "FERIA" in val or "FOLGA" in val: return "ABONADO"
    if "DESLIG" in val or "SAIU" in val: return "AFASTADO"
    return "ACIONADO"


def parse_time(value: object) -> str | None:
    if isinstance(value, dt.time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, dt.datetime):
        return value.strftime("%H:%M:%S")
    if isinstance(value, (int, float)) and 0 <= value < 1:
        seconds = round(value * 86400) % 86400
        return f"{seconds // 3600:02d}:{seconds % 3600 // 60:02d}:00"
    if isinstance(value, str):
        val = value.strip()
        if re.fullmatch(r"\d{1,2}:\d{2}(?::\d{2})?", val):
            parts = val.split(":")
            return f"{int(parts[0]):02d}:{int(parts[1]):02d}:{int(parts[2]) if len(parts) > 2 else 0:02d}"
    return None


def calculate_minutes(entry_time_str: str | None, exit_time_str: str | None) -> int | None:
    if not entry_time_str or not exit_time_str:
        return None
    try:
        e_h, e_m, e_s = map(int, entry_time_str.split(":"))
        x_h, x_m, x_s = map(int, exit_time_str.split(":"))
        start_min = e_h * 60 + e_m
        end_min = x_h * 60 + x_m
        diff = end_min - start_min
        if diff < 0:
            diff += 24 * 60
        return diff
    except Exception:
        return None


def norm_header(txt: object) -> str:
    if not txt:
        return ""
    s = unicodedata.normalize("NFKD", str(txt)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]", "", s)


def load_env() -> dict[str, str]:
    result: dict[str, str] = {}
    for line in (ROOT / ".env.local").read_text(encoding="utf-8").splitlines():
        if line and not line.lstrip().startswith("#") and "=" in line:
            key, value = line.split("=", 1)
            result[key.strip()] = value.strip().strip('"').strip("'")
    return result


def request_json(url: str, key: str, method: str = "GET", body: object | None = None, prefer: str | None = None):
    headers = {"apikey": key, "Authorization": f"Bearer {key}"}
    data = None
    if body is not None:
        headers["Content-Type"] = "application/json"
        data = json.dumps(body, ensure_ascii=False).encode("utf-8")
    if prefer:
        headers["Prefer"] = prefer
    request = urllib.request.Request(url, data=data, headers=headers, method=method)
    try:
        with urllib.request.urlopen(request, timeout=60) as response:
            raw = response.read()
            return json.loads(raw) if raw else None
    except urllib.error.HTTPError as error:
        detail = error.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"Supabase retornou HTTP {error.code}: {detail}") from error


def get_or_create_location(name: str, region: str, location_by_key: dict, base_url: str, service_key: str, apply: bool) -> str | None:
    k = match_key(name)
    if not k:
        return None
    if k in location_by_key:
        return location_by_key[k]
    if not apply:
        return None
    res = request_json(
        f"{base_url}/locations",
        service_key,
        method="POST",
        body={"name": name.strip().upper(), "region": region, "is_active": True},
        prefer="resolution=merge-duplicates,return=representation",
    )
    if res:
        loc_id = res[0]["id"]
        location_by_key[k] = loc_id
        return loc_id
    return None


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--apply", action="store_true", help="Grava no Supabase.")
    parser.add_argument("--limit-sheets", type=int, default=None, help="Limita o numero de abas processadas.")
    args = parser.parse_args()

    env = load_env()
    base_url = env["NEXT_PUBLIC_SUPABASE_URL"].rstrip("/") + "/rest/v1"
    service_key = env["SUPABASE_SERVICE_ROLE_KEY"]

    # Carregar entidades do Supabase
    db_drivers = request_json(f"{base_url}/drivers?select=id,name,login_matricula", service_key)
    driver_by_key = {match_key(d["name"]): d["id"] for d in db_drivers if d.get("name")}
    driver_by_login = {d["login_matricula"]: d["id"] for d in db_drivers if d.get("login_matricula")}

    db_locations = request_json(f"{base_url}/locations?select=id,name,region", service_key)
    location_by_key = {match_key(l["name"]): l["id"] for l in db_locations if l.get("name")}

    db_vehicles = request_json(f"{base_url}/vehicles?select=id,plate_code", service_key)
    vehicle_by_plate = {v["plate_code"].upper(): v["id"] for v in db_vehicles if v.get("plate_code")}

    workbook = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)

    daily_sheets_meta = []
    for sheet_name in workbook.sheetnames:
        m = re.match(r"^(\d{2})\.(\d{2})\.(\d{2})(?:\s*-\s*UBJP)?$", sheet_name.strip())
        if m:
            day, month, year = int(m.group(1)), int(m.group(2)), 2000 + int(m.group(3))
            is_ubjp = "UBJP" in sheet_name
            region = "UBJP" if is_ubjp else "SP"
            date_str = f"{year:04d}-{month:02d}-{day:02d}"
            daily_sheets_meta.append((sheet_name, date_str, region))

    print(f"Total de {len(daily_sheets_meta)} abas diarias encontradas.")
    if args.limit_sheets:
        daily_sheets_meta = daily_sheets_meta[:args.limit_sheets]
        print(f"Limitando para {len(daily_sheets_meta)} abas.")

    total_dispatches = 0

    for sheet_name, date_str, region in daily_sheets_meta:
        sheet = workbook[sheet_name]

        header_map = {}
        for r in range(1, 6):
            row_vals = [cell for cell in sheet.iter_rows(min_row=r, max_row=r, values_only=True)][0]
            for idx, col in enumerate(row_vals):
                k = norm_header(col)
                if not k:
                    continue
                if "condutor" in k or "colaborador" in k: header_map["name"] = idx
                elif "login" in k: header_map["login"] = idx
                elif "horario" in k or "hora" in k: header_map["horario"] = idx
                elif "prevprox" in k or "prevproxdia" in k: header_map["prev_prox"] = idx
                elif "previsao" in k: header_map["previsao"] = idx
                elif "acionado" in k: header_map["status"] = idx
                elif "motivo" in k: header_map["abono"] = idx
                elif "local" in k: header_map["local"] = idx
                elif "entrada" in k: header_map["entrada"] = idx
                elif "frota" in k: header_map["frota"] = idx
                elif "cv" in k or "vc" in k: header_map["cargo"] = idx
                elif "destinofinal" in k: header_map["destino_final"] = idx
                elif "destino" in k and "destinofinal" not in k: header_map["destino"] = idx
                elif "saida" in k: header_map["saida"] = idx
                elif "observa" in k: header_map["obs"] = idx
                elif "diaesc" in k or "diadeesc" in k: header_map["dia_esc"] = idx
            if "name" in header_map and "entrada" in header_map:
                break

        if "name" not in header_map:
            print(f"Aba {sheet_name}: cabecalho nao reconhecido, pulando.")
            continue

        dispatches_dict = {}
        for row in sheet.iter_rows(min_row=5, values_only=True):
            if not row or not any(row):
                continue
            raw_name = clean_str(row[header_map["name"]]) if "name" in header_map and header_map["name"] < len(row) else None
            if not raw_name:
                continue
            name = re.sub(r"\s+-\s+VOLTA\s+\d{1,2}/\d{1,2}\s*$", "", raw_name, flags=re.IGNORECASE)
            key = match_key(name)
            login = clean_str(row[header_map["login"]]) if "login" in header_map and header_map["login"] < len(row) else None

            driver_id = driver_by_key.get(key) or (driver_by_login.get(login) if login else None)
            if not driver_id:
                continue

            # Local / Origem
            origin_name = clean_str(row[header_map["local"]]) if "local" in header_map and header_map["local"] < len(row) else None
            origin_id = get_or_create_location(origin_name, region, location_by_key, base_url, service_key, args.apply) if origin_name else None

            # Destino
            dest_name = clean_str(row[header_map["destino"]]) if "destino" in header_map and header_map["destino"] < len(row) else None
            dest_id = get_or_create_location(dest_name, region, location_by_key, base_url, service_key, args.apply) if dest_name else None

            # Destino Final
            dest_final_name = clean_str(row[header_map["destino_final"]]) if "destino_final" in header_map and header_map["destino_final"] < len(row) else None
            dest_final_id = get_or_create_location(dest_final_name, region, location_by_key, base_url, service_key, args.apply) if dest_final_name else dest_id

            # Frota
            fleet_code = clean_str(row[header_map["frota"]]) if "frota" in header_map and header_map["frota"] < len(row) else None
            vehicle_id = None
            if fleet_code:
                norm_fleet = fleet_code.upper()
                if norm_fleet in vehicle_by_plate:
                    vehicle_id = vehicle_by_plate[norm_fleet]
                elif args.apply:
                    v_res = request_json(
                        f"{base_url}/vehicles",
                        service_key,
                        method="POST",
                        body={"plate_code": norm_fleet, "vehicle_type": "CAVALO_MECANICO", "is_active": True},
                        prefer="resolution=merge-duplicates,return=representation",
                    )
                    if v_res:
                        vehicle_id = v_res[0]["id"]
                        vehicle_by_plate[norm_fleet] = vehicle_id

            entry_t = parse_time(row[header_map["entrada"]]) if "entrada" in header_map and header_map["entrada"] < len(row) else None
            exit_t = parse_time(row[header_map["saida"]]) if "saida" in header_map and header_map["saida"] < len(row) else None
            forecast_t = parse_time(row[header_map["previsao"]]) if "previsao" in header_map and header_map["previsao"] < len(row) else None
            next_forecast_t = parse_time(row[header_map["prev_prox"]]) if "prev_prox" in header_map and header_map["prev_prox"] < len(row) else None
            
            raw_status = clean_str(row[header_map["status"]]) if "status" in header_map and header_map["status"] < len(row) else None
            status_val = normalize_status(raw_status)

            abono_reason = clean_str(row[header_map["abono"]]) if "abono" in header_map and header_map["abono"] < len(row) else None
            cargo_status = clean_str(row[header_map["cargo"]]) if "cargo" in header_map and header_map["cargo"] < len(row) else None
            if cargo_status:
                cargo_status = cargo_status.upper()
                if cargo_status not in ("C", "V"):
                    cargo_status = None
            obs = clean_str(row[header_map["obs"]]) if "obs" in header_map and header_map["obs"] < len(row) else None
            dia_esc = clean_str(row[header_map["dia_esc"]]) if "dia_esc" in header_map and header_map["dia_esc"] < len(row) else None
            if dia_esc and dia_esc.endswith(".0"):
                dia_esc = dia_esc[:-2]

            worked_min = calculate_minutes(entry_t, exit_t)

            item = {
                "driver_id": driver_id,
                "forecast_start_time": forecast_t,
                "dispatch_status": status_val,
                "abono_reason": abono_reason,
                "origin_location_id": origin_id,
                "entry_time": entry_t,
                "vehicle_id": vehicle_id,
                "cargo_status": cargo_status,
                "destination_location_id": dest_id,
                "final_destination_location_id": dest_final_id,
                "exit_time": exit_t,
                "observations": obs,
                "scale_day_status": dia_esc,
                "next_day_forecast_time": next_forecast_t,
                "total_worked_minutes": worked_min,
            }

            if driver_id not in dispatches_dict or entry_t or exit_t:
                dispatches_dict[driver_id] = item

        dispatches = list(dispatches_dict.values())

        if args.apply and dispatches:
            sheet_res = request_json(
                f"{base_url}/daily_sheets?on_conflict=date,region",
                service_key,
                method="POST",
                body=[{"date": date_str, "region": region, "status": "CLOSED" if date_str < "2026-09-01" else "OPEN"}],
                prefer="resolution=merge-duplicates,return=representation",
            )
            daily_sheet_id = sheet_res[0]["id"]

            for item in dispatches:
                item["daily_sheet_id"] = daily_sheet_id

            for i in range(0, len(dispatches), 100):
                batch = dispatches[i:i + 100]
                request_json(
                    f"{base_url}/dispatches?on_conflict=daily_sheet_id,driver_id",
                    service_key,
                    method="POST",
                    body=batch,
                    prefer="resolution=merge-duplicates",
                )

        total_dispatches += len(dispatches)
        print(f"[{date_str} - {region}] {len(dispatches)} despachos sincronizados.")

    print(f"\nFinalizado: {len(daily_sheets_meta)} abas processadas, {total_dispatches} despachos mapeados.")


if __name__ == "__main__":
    main()
