"""Sincroniza as escalas mensais e registros diários com o Supabase.

Processa as abas de Escala Mensal do Excel TESTE.xlsx e grava em:
- monthly_schedules
- monthly_schedule_entries
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "referencias" / "TESTE.xlsx"

MONTH_MAP = {
    "Junho": 6,
    "Julho": 7,
    "Agosto": 8,
    "Setembro": 9,
}

SCHEDULE_SHEETS = [
    ("Escala Junho 2026 - SP", 6, 2026, "SP"),
    ("Escala Junho 2026 - MG", 6, 2026, "MG"),
    ("Escala Julho 2026 - SP", 7, 2026, "SP"),
    ("Escala Julho 2026 - MG", 7, 2026, "MG"),
    ("Escala Agosto 2026 - SP", 8, 2026, "SP"),
    ("Escala Agosto 2026 - MG", 8, 2026, "MG"),
    ("Escala Setembro 2026 - SP", 9, 2026, "SP"),
    ("Escala Setembro 2026 - MG", 9, 2026, "MG"),
]


def normalize_name(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def match_key(value: object) -> str:
    text = unicodedata.normalize("NFKD", normalize_name(value)).encode("ascii", "ignore").decode().upper()
    return re.sub(r"[^A-Z0-9]", "", text)


def clean_login(value: object) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def parse_time(value: object) -> str | None:
    if isinstance(value, dt.time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, (int, float)) and 0 <= value < 1:
        seconds = round(value * 86400) % 86400
        return f"{seconds // 3600:02d}:{seconds % 3600 // 60:02d}:00"
    if isinstance(value, str) and re.fullmatch(r"\d{1,2}:\d{2}(?::\d{2})?", value.strip()):
        parts = value.strip().split(":")
        return f"{int(parts[0]):02d}:{int(parts[1]):02d}:{int(parts[2]) if len(parts) > 2 else 0:02d}"
    return None


def row_metadata(values: tuple[object, ...], first_meta_column: int) -> tuple[str | None, str | None]:
    metadata = values[first_meta_column - 1:]
    start_time = next((parsed for value in metadata if (parsed := parse_time(value))), None)
    shift = next((str(value).strip().upper() for value in metadata if isinstance(value, str) and re.fullmatch(r"[A-Z]", value.strip().upper())), None)
    return shift, start_time


def load_env() -> dict[str, str]:
    result: dict[str, str] = {}
    for line in (ROOT / ".env.local").read_text(encoding="utf-8").splitlines():
        if line and not line.lstrip().startswith("#") and "=" in line:
            key, value = line.split("=", 1)
            result[key.strip()] = value.strip().strip('"').strip("'")
    return result


def request_json(url: str, key: str, method: str = "GET", body: object | None = None, prefer: str | None = None):
    headers = {"apikey": key, "Authorization": f"Bearer {key}"}
    data = None
    if body is not None:
        headers["Content-Type"] = "application/json"
        data = json.dumps(body, ensure_ascii=False).encode("utf-8")
    if prefer:
        headers["Prefer"] = prefer
    request = urllib.request.Request(url, data=data, headers=headers, method=method)
    try:
        with urllib.request.urlopen(request, timeout=60) as response:
            raw = response.read()
            return json.loads(raw) if raw else None
    except urllib.error.HTTPError as error:
        detail = error.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"Supabase retornou HTTP {error.code}: {detail}") from error


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--apply", action="store_true", help="Grava no Supabase.")
    parser.add_argument("--only-september", action="store_true", help="Processa apenas Setembro/2026.")
    args = parser.parse_args()

    env = load_env()
    base_url = env["NEXT_PUBLIC_SUPABASE_URL"].rstrip("/") + "/rest/v1"
    service_key = env["SUPABASE_SERVICE_ROLE_KEY"]

    # Carregar condutores do Supabase para fazer o mapeamento
    db_drivers = request_json(f"{base_url}/drivers?select=id,name,login_matricula,is_active", service_key)
    driver_by_key = {match_key(d["name"]): d["id"] for d in db_drivers if d.get("name")}
    driver_by_login = {d["login_matricula"]: d["id"] for d in db_drivers if d.get("login_matricula")}

    workbook = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    sheets_to_process = [s for s in SCHEDULE_SHEETS if not args.only_september or s[1] == 9]

    total_schedules = 0
    total_entries = 0

    for sheet_name, month, year, region in sheets_to_process:
        if sheet_name not in workbook.sheetnames:
            print(f"Aba {sheet_name} não encontrada, pulando...")
            continue

        sheet = workbook[sheet_name]
        is_mg = region == "MG"
        days_in_month = 31 if month in (1, 3, 5, 7, 8, 10, 12) else 30
        first_day_col = 3 if is_mg else 2 # 0-indexed: Col D (index 3) ou Col C (index 2)
        login_col = 2 if is_mg else 1
        meta_col = first_day_col + days_in_month + 1

        print(f"\nProcessando {sheet_name} ({month}/{year} - {region})...")

        entries_by_cell = {}
        unmatched_rows = 0

        for row_idx, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            if not values or not values[0]:
                continue

            raw_name = normalize_name(values[0])
            name = re.sub(r"\s+-\s+VOLTA\s+\d{1,2}/\d{1,2}\s*$", "", raw_name, flags=re.IGNORECASE)
            key = match_key(name)
            login = clean_login(values[login_col]) if len(values) > login_col else None

            # Prioriza a chave do nome para evitar colisão de logins legados corrigidos
            driver_id = driver_by_key.get(key) or (driver_by_login.get(login) if login else None)
            if not driver_id:
                unmatched_rows += 1
                continue

            shift, start_time = row_metadata(values, meta_col)

            for day in range(1, days_in_month + 1):
                col_idx = first_day_col + (day - 1)
                if col_idx < len(values):
                    val = values[col_idx]
                    if val is not None and str(val).strip() != "":
                        code_str = str(val).strip().upper()
                        if code_str.endswith(".0"):
                            code_str = code_str[:-2]
                        cell_k = (driver_id, day)
                        entries_by_cell[cell_k] = {
                            "driver_id": driver_id,
                            "day_of_month": day,
                            "scale_code": code_str,
                            "shift_letter": shift,
                            "shift_time": start_time,
                        }

        schedule_entries = list(entries_by_cell.values())
        print(f"  -> {len(schedule_entries)} registros unicos encontrados ({unmatched_rows} linhas nao mapeadas).")

        if args.apply:
            # 1. Upsert monthly_schedule
            sched_body = {
                "month": month,
                "year": year,
                "region": region,
                "status": "ACTIVE",
            }
            sched_res = request_json(
                f"{base_url}/monthly_schedules?on_conflict=month,year,region",
                service_key,
                method="POST",
                body=[sched_body],
                prefer="resolution=merge-duplicates,return=representation",
            )
            schedule_id = sched_res[0]["id"]

            # 2. Upsert entries em lotes de 500
            for i in range(0, len(schedule_entries), 500):
                batch = schedule_entries[i:i + 500]
                for item in batch:
                    item["schedule_id"] = schedule_id
                request_json(
                    f"{base_url}/monthly_schedule_entries?on_conflict=schedule_id,driver_id,day_of_month",
                    service_key,
                    method="POST",
                    body=batch,
                    prefer="resolution=merge-duplicates,return=representation",
                )
            print(f"  [OK] Gravado no Supabase (Schedule ID: {schedule_id}).")

        total_schedules += 1
        total_entries += len(schedule_entries)

    print(f"\nResumo: {total_schedules} escalas mensais processadas, {total_entries} entradas.")
    if not args.apply:
        print("Modo de validacao. Execute com --apply para gravar.")


if __name__ == "__main__":
    main()
