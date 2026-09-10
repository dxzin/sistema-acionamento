import openpyxl

wb = openpyxl.load_workbook('referencias/TESTE.xlsx', data_only=False)

def search_abonar_columns():
    print("=== SEARCHING ABONAR COLUMNS ACROSS ALL UBJP SHEETS ===")
    found = 0
    for name in wb.sheetnames:
        if 'UBJP' in name and 'Escala' not in name:
            ws = wb[name]
            if not hasattr(ws, 'max_row') or ws.max_row < 5:
                continue
            for r in range(4, ws.max_row+1):
                for c in range(1, ws.max_column+1):
                    h = str(ws.cell(4, c).value)
                    if 'ABONAR' in h.upper():
                        val = ws.cell(r, c).value
                        if val is not None and r > 4:
                            print(f"Sheet {name}, Row {r}, Col {c} ({h}): {val}")
                            found += 1
                            if found > 20:
                                return
    print(f"Total non-empty abonar cells found: {found}")

search_abonar_columns()
