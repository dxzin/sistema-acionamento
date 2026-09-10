import openpyxl, re

wb = openpyxl.load_workbook('referencias/TESTE.xlsx', data_only=False)

table_to_sheet = {}
for name in wb.sheetnames:
    ws = wb[name]
    if hasattr(ws, "tables"):
        for t in ws.tables.values():
            table_to_sheet[t.name] = name

print(f"Total tables: {len(table_to_sheet)}")

def resolve_formula(f, current_sheet):
    if not isinstance(f, str):
        return str(f)
    def repl(match):
        t_name = match.group(0)
        sheet = table_to_sheet.get(t_name, "UNKNOWN_TABLE")
        if sheet == current_sheet:
            return "THIS_DAY"
        return f"SHEET('{sheet}')"
    return re.sub(r'Tabela[0-9a-zA-Z_]+', repl, f)

print("=== CHECKING FORMULA RESOLUTION FOR 02.06.26 ===")
ws = wb['02.06.26']
for c in range(1, ws.max_column+1):
    h = ws.cell(4, c).value
    col_letter = openpyxl.utils.get_column_letter(c)
    f5 = resolve_formula(ws.cell(5, c).value, '02.06.26')
    f6 = resolve_formula(ws.cell(6, c).value, '02.06.26')
    print(f"Col {col_letter:2s} ({c:2d}) [{h}]:")
    print(f"    R5: {f5}")
    if f6 != f5 and f6 != "None":
        print(f"    R6: {f6}")

print("\n=== CHECKING FORMULA RESOLUTION FOR 02.06.26  - UBJP ===")
ws_u = wb['02.06.26  - UBJP']
for c in range(1, ws_u.max_column+1):
    h = ws_u.cell(4, c).value
    col_letter = openpyxl.utils.get_column_letter(c)
    f5 = resolve_formula(ws_u.cell(5, c).value, '02.06.26  - UBJP')
    f6 = resolve_formula(ws_u.cell(6, c).value, '02.06.26  - UBJP')
    print(f"Col {col_letter:2s} ({c:2d}) [{h}]:")
    print(f"    R5: {f5}")
    if f6 != f5 and f6 != "None":
        print(f"    R6: {f6}")

