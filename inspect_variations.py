import openpyxl, re

wb = openpyxl.load_workbook('referencias/TESTE.xlsx', data_only=False)

def inspect_all_sheet_variations():
    print("=== CHECKING HEADERS ACROSS ALL DAILY GENERAL SHEETS ===")
    gen_headers = {}
    for name in wb.sheetnames:
        if 'UBJP' not in name and 'Escala' not in name and 'REFER' not in name and name not in ['Planilha1', 'Grfico3']:
            ws = wb[name]
            if not hasattr(ws, 'max_row') or ws.max_row < 4:
                continue
            headers = tuple(ws.cell(4, c).value for c in range(1, ws.max_column+1) if ws.cell(4, c).value is not None)
            gen_headers.setdefault(headers, []).append(name)
    
    for h, names in gen_headers.items():
        print(f"\nHeader variant ({len(h)} cols) in {len(names)} sheets (e.g. {names[:3]}):")
        for idx, col_h in enumerate(h, 1):
            print(f"  {idx:2d}: {col_h}")

    print("\n=== CHECKING HEADERS ACROSS ALL DAILY UBJP SHEETS ===")
    ubjp_headers = {}
    for name in wb.sheetnames:
        if 'UBJP' in name and 'Escala' not in name and 'REFER' not in name:
            ws = wb[name]
            if not hasattr(ws, 'max_row') or ws.max_row < 4:
                continue
            headers = tuple(ws.cell(4, c).value for c in range(1, ws.max_column+1) if ws.cell(4, c).value is not None)
            ubjp_headers.setdefault(headers, []).append(name)
    
    for h, names in ubjp_headers.items():
        print(f"\nUBJP Header variant ({len(h)} cols) in {len(names)} sheets (e.g. {names[:3]}):")
        for idx, col_h in enumerate(h, 1):
            print(f"  {idx:2d}: {col_h}")

inspect_all_sheet_variations()
