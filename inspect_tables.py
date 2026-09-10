import openpyxl, re

wb = openpyxl.load_workbook('referencias/TESTE.xlsx', data_only=False)

def get_table_sheet_map():
    mapping = {}
    for name in wb.sheetnames:
        ws = wb[name]
        if hasattr(ws, "tables"):
            for t in ws.tables.values():
                mapping[t.name] = (name, t.ref)
    return mapping

t_map = get_table_sheet_map()
print(f"Total tables mapped: {len(t_map)}")

for s_name in ['01.06.26', '02.06.26', '03.06.26', '01.07.26', '02.07.26']:
    ws = wb[s_name]
    print(f"\n==================== SHEET: {s_name} ====================")
    this_tables = [t.name for t in ws.tables.values()] if hasattr(ws, 'tables') else []
    print(f"Tables in this sheet: {this_tables}")
    for col_idx in [4, 7, 15, 16, 17, 18, 19]:
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        header = ws.cell(4, col_idx).value
        val5 = ws.cell(5, col_idx).value
        val6 = ws.cell(6, col_idx).value
        print(f"Col {col_letter} ({header}):")
        print(f"   R5 formula: {val5}")
        if val6 != val5:
            print(f"   R6 formula: {val6}")

