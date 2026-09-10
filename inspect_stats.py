import openpyxl

wb = openpyxl.load_workbook('referencias/TESTE.xlsx', data_only=True)

all_drivers = set()
all_fleets = set()
all_destinations = set()
all_reasons_abono = set()
all_approvers = set()
all_bases = set()

for name in wb.sheetnames:
    if 'Escala' in name:
        ws = wb[name]
        for r in range(2, ws.max_row+1):
            driver = ws.cell(r, 1).value
            if driver:
                all_drivers.add(str(driver).strip())
            if 'MG' in name or 'UBJP' in name:
                base = ws.cell(r, 2).value
                if base:
                    all_bases.add(str(base).strip())
    elif any(char.isdigit() for char in name) and '.' in name:
        ws = wb[name]
        if ws.max_row >= 5:
            # check headers
            headers = [ws.cell(4, c).value for c in range(1, ws.max_column+1)]
            for r in range(5, ws.max_row+1):
                for c, h in enumerate(headers, 1):
                    val = ws.cell(r, c).value
                    if val is not None:
                        val_str = str(val).strip()
                        h_upper = str(h).upper()
                        if 'FROTA' in h_upper:
                            all_fleets.add(val_str)
                        elif 'DESTINO' in h_upper:
                            all_destinations.add(val_str)
                        elif 'MOTIVO' in h_upper and 'ABONO' in h_upper:
                            all_reasons_abono.add(val_str)
                        elif 'AUTORIZA' in h_upper and 'ABONO' in h_upper:
                            all_approvers.add(val_str)

print(f"Total Unique Drivers: {len(all_drivers)}")
print(f"Total Unique Fleets: {len(all_fleets)}")
print(f"Bases identified: {sorted(list(all_bases))}")
print(f"Destinations identified: {sorted(list(all_destinations))}")
print(f"Abono reasons identified: {sorted(list(all_reasons_abono))}")
print(f"Abono approvers/values identified: {sorted(list(all_approvers))}")

